from flask import Flask, request, jsonify
from datetime import datetime
from flask_cors import CORS
from flask import session
from werkzeug.security import check_password_hash
from flask import send_file
from io import BytesIO
import pandas as pd
import json
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import or_
from sqlalchemy.exc import IntegrityError
from models import (
    db,
    User,
    Quote,
    LineItem,
    Product,
    NotesLibrary,
    LineItemNote,
    Company,
    Contact,
    generate_unique_quote_number,
)
from config import Config

app = Flask(__name__)
app.config.from_object(Config)
import os

app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="None",
    # SESSION_COOKIE_SECURE=True,
    SESSION_COOKIE_SECURE=os.environ.get("FLASK_ENV") != "development",
)


CORS(app, supports_credentials=True, origins=["https://quotiq.onrender.com", "http://localhost:5173", "http://127.0.0.1:5173"])
db.init_app(app)

with app.app_context():
    db.create_all()


# ---------------------------
# Helpers
# ---------------------------
def money(value):
    return float(value or 0)


def num(value, default=0):
    if value is None or value == "":
        return default
    return float(value)


def bool_value(value, default=False):
    if value is None:
        return default
    return bool(value)


def round_to_preferred(value):
    return round(value)


def calculate_line_item(item):
    qty = float(item.qty or 1)
    list_price = float(item.list_price or 0)
    multiplier = float(item.multiplier or 1)
    markup = float(item.markup or 0)
    freight = float(item.freight or 0)
    startup = float(item.startup or 0)
    surcharge = float(item.surcharge or 0)

    net_cost = list_price * (1 + surcharge) * multiplier
    raw_sell = (net_cost * (1 + markup)) + freight + startup
    sell_price = round_to_preferred(raw_sell)
    total_price = sell_price * qty

    item.net_cost = round(net_cost, 2)
    item.sell_price = round(sell_price, 2)
    item.total_price = round(total_price, 2)


def line_item_note_to_dict(note):
    return {
        "id": note.id,
        "line_item_id": note.line_item_id,
        "note_library_id": note.note_library_id,
        "category": note.category,
        "label": note.label,
        "text": note.text,
        "is_custom": note.is_custom,
        "is_selected": note.is_selected,
        "sort_order": note.sort_order,
    }


def line_item_to_dict(item):
    return {
        "id": item.id,
        "quote_id": item.quote_id,
        "sort_order": item.sort_order,
        "tag": item.tag,
        "vendor": item.vendor,
        "qty": item.qty,
        "description": item.description,
        "item": item.item,
        "type": item.type,
        "series": item.series,
        "model": item.model,
        "part_number": item.part_number,
        "list_price": money(item.list_price),
        "multiplier": money(item.multiplier),
        "markup": money(item.markup),
        "freight": money(item.freight),
        "startup": money(item.startup),
        "surcharge": money(item.surcharge),
        "net_cost": money(item.net_cost),
        "sell_price": money(item.sell_price),
        "total_price": money(item.total_price),
        "terms": item.terms,
        "notes": item.notes,
        "included": item.included,
        "notes_selected": [
            line_item_note_to_dict(n)
            for n in sorted(item.notes_selected, key=lambda n: n.sort_order or 0)
        ],
    }


def quote_to_dict(q):
    sorted_line_items = sorted(q.line_items, key=lambda item: item.sort_order or 0)

    return {
        "id": q.id,
        # "date": q.date.strftime("%m/%d/%Y") if q.date else "",
        "date": q.date.isoformat() if q.date else "",
        "quote_number": q.quote_number,
        "bid_date": q.bid_date,
        "contact": q.contact or [],
        "project": q.project,
        "to_company": q.to_company,
        "attention": q.attention,
        "location": q.location,
        "notes": q.notes,
        "status": q.status,
        "freight_terms": q.freight_terms or "FOB",
        "created_at": q.created_at.strftime("%m/%d/%Y") if q.created_at else "",
        "updated_at": q.updated_at.strftime("%m/%d/%Y") if q.updated_at else "",
        "line_items": [line_item_to_dict(item) for item in sorted_line_items],
        "total": sum(float(item.total_price or 0) for item in sorted_line_items),
        "created_by": q.created_by,
        "locked_by": q.locked_by,
        "locked_at": q.locked_at.isoformat() if q.locked_at else None,
    }


def product_to_dict(p):
    return {
        "id": p.id,
        "name": p.name,
        "tag": p.tag,
        "vendor": p.vendor,
        "manufacturer": p.manufacturer,
        "category": p.category,
        "type": p.type,
        "series": p.series,
        "model": p.model,
        "part_number": p.part_number,
        "description": p.description,
        "list_price": money(p.list_price),
        "surcharge": money(p.surcharge),
        "multiplier": money(p.multiplier),
        "net_cost": money(p.net_cost),
        "notes": p.notes,
        "is_active": p.is_active,
    }


def note_to_dict(n):
    return {
        "id": n.id,
        "item": n.item,
        "type": n.type,
        "category": n.category,
        "series": n.series,
        "model": n.model,
        "text": n.text,
        "note_type": n.note_type,
        "default_selected": n.default_selected,
        "sort_order": n.sort_order,
        "is_active": n.is_active,
    }


def company_to_dict(c, include_contacts=True):
    data = {
        "id": c.id,
        "name": c.name,
        "type": c.type,
        "address1": c.address1,
        "address2": c.address2,
        "city": c.city,
        "state": c.state,
        "zipcode": c.zipcode,
        "notes": c.notes,
        "is_active": c.is_active,
        "website": c.website,
        "account_number": c.account_number,
        "tax_id": c.tax_id,
        "payment_terms": c.payment_terms,
    }
    if include_contacts:
        data["contacts"] = [contact_to_dict(x) for x in c.contacts]
    return data


def contact_to_dict(c):
    return {
        "id": c.id,
        "company_id": c.company_id,
        "company_name": c.company.name if c.company else "",
        "first_name": c.first_name,
        "last_name": c.last_name,
        "role": c.role,
        "email": c.email,
        "tel": c.tel,
        "mobile": c.mobile,
        "notes": c.notes,
        "is_active": c.is_active,
    }


def apply_product_to_line_item(item, product):
    item.item = product.category
    item.type = product.type
    item.series = product.series
    item.model = product.model
    item.part_number = product.part_number
    item.vendor = product.vendor
    item.description = product.description
    item.list_price = product.list_price or 0
    item.multiplier = product.multiplier or 1
    item.surcharge = product.surcharge or 0


# ---------------------------
# Quotes
# ---------------------------
@app.route("/quotes", methods=["GET"])
def get_quotes():
    search = request.args.get("search", "").strip()
    line_search = request.args.get("line_search", "").strip()
    status = request.args.get("status", "").strip()
    customer = request.args.get("customer", "").strip()
    location = request.args.get("location", "").strip()

    salesman = request.args.get("salesman", "").strip()

    quote_date_from = request.args.get("quote_date_from", "").strip()
    quote_date_to = request.args.get("quote_date_to", "").strip()
    bid_date_from = request.args.get("bid_date_from", "").strip()
    bid_date_to = request.args.get("bid_date_to", "").strip()

    sort_by = request.args.get("sort_by", "id")
    direction = request.args.get("direction", "desc")

    query = Quote.query

    if search:
        like = f"%{search}%"
        query = query.filter(
            or_(
                Quote.quote_number.ilike(like),
                Quote.project.ilike(like),
                Quote.attention.ilike(like),
            )
        )

    if line_search:
        like = f"%{line_search}%"
        query = query.join(LineItem).filter(LineItem.description.ilike(like))

    if status:
        query = query.filter(Quote.status == status)

    if customer:
        query = query.filter(Quote.to_company.ilike(f"%{customer}%"))

    if location:
        query = query.filter(Quote.location.ilike(f"%{location}%"))

    if quote_date_from:
        query = query.filter(Quote.date >= datetime.strptime(quote_date_from, "%Y-%m-%d").date())

    if quote_date_to:
        query = query.filter(Quote.date <= datetime.strptime(quote_date_to, "%Y-%m-%d").date())

    if bid_date_from:
        query = query.filter(Quote.bid_date >= bid_date_from)

    if bid_date_to:
        query = query.filter(Quote.bid_date <= bid_date_to)

    if salesman:
        query = query.filter(Quote.contact.cast(db.String).ilike(f"%{salesman}%"))

    sort_map = {
        "id": Quote.id,
        "quote_number": Quote.quote_number,
        "date": Quote.date,
        "bid_date": Quote.bid_date,
        "created_at": Quote.created_at,
        "status": Quote.status,
        "project": Quote.project,
        "to_company": Quote.to_company,
        "attention": Quote.attention,
        "location": Quote.location,
    }

    col = sort_map.get(sort_by, Quote.id)
    query = query.order_by(col.asc() if direction == "asc" else col.desc())

    return jsonify([quote_to_dict(q) for q in query.distinct().all()])


@app.route("/quotes/<int:id>", methods=["GET"])
def get_quote(id):
    return jsonify(quote_to_dict(Quote.query.get_or_404(id)))


@app.route("/quotes/<int:id>/lock", methods=["POST"])
def lock_quote(id):
    if "user_id" not in session:
        return jsonify({"error": "Not logged in"}), 401

    quote = Quote.query.get_or_404(id)
    user_name = session["user_name"]

    if quote.locked_by and quote.locked_by != user_name:
        return jsonify({
            "error": f"This quote is currently being edited by {quote.locked_by}."
        }), 409

    quote.locked_by = user_name
    quote.locked_at = datetime.utcnow()

    db.session.commit()

    return jsonify(quote_to_dict(quote))


@app.route("/quotes/<int:id>/unlock", methods=["POST"])
def unlock_quote(id):
    if "user_id" not in session:
        return jsonify({"error": "Not logged in"}), 401

    quote = Quote.query.get_or_404(id)

    quote.locked_by = None
    quote.locked_at = None

    db.session.commit()

    return jsonify(quote_to_dict(quote))


@app.route("/quotes", methods=["POST"])
def create_quote():
    data = request.json or {}
    quote_number = str(data.get("quote_number") or "").strip()
    if not quote_number:
        return jsonify({"error": "Enter a quote number before saving."}), 400
    existing = Quote.query.filter_by(quote_number=quote_number).first()
    if existing:
        return jsonify({
            "error": f"Quote # {quote_number} already exists. Enter a different quote number."
        }), 409
    quote = Quote(
        quote_number=quote_number,
        date=datetime.strptime(data.get("date"),"%Y-%m-%d").date() if data.get("date") else datetime.utcnow().date(),
        bid_date=data.get("bid_date", "N/A"),
        contact=data.get("contact", []),
        project=data.get("project"),
        to_company=data.get("to_company"),
        attention=data.get("attention"),
        location=data.get("location"),
        freight_terms=data.get("freight_terms", "FOB"),
        status=data.get("status", "Not Started"),
        notes=data.get("notes"),
        created_by=session.get("user_name"),
    )
    db.session.add(quote)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({
            "error": f"Quote # {quote_number} already exists. Enter a different quote number."
        }), 409
    return jsonify(quote_to_dict(quote)), 201


@app.route("/quotes/<int:id>", methods=["PUT"])
def update_quote(id):
    quote = Quote.query.get_or_404(id)
    data = request.json or {}
    # quote date
    if data.get("date"):
        quote.date = datetime.strptime(data.get("date"), "%Y-%m-%d").date()

    # bid due date
    quote.bid_date = data.get("bid_date", quote.bid_date)
    quote.contact = data.get("contact", quote.contact or [])
    quote.project = data.get("project", quote.project)
    quote.to_company = data.get("to_company", quote.to_company)
    quote.attention = data.get("attention", quote.attention)
    quote.location = data.get("location", quote.location)
    quote.freight_terms = data.get("freight_terms", quote.freight_terms or "FOB")
    quote.status = data.get("status", quote.status)
    quote.notes = data.get("notes", quote.notes)
    if data.get("quote_number") is not None:
        incoming = str(data.get("quote_number") or "").strip()
        if not incoming:
            return jsonify({"error": "Enter a quote number before saving."}), 400
        taken = Quote.query.filter(
            Quote.quote_number == incoming,
            Quote.id != quote.id,
        ).first()
        if taken:
            return jsonify({
                "error": f"Quote # {incoming} already exists. Enter a different quote number."
            }), 409
        quote.quote_number = incoming

    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify({
            "error": f"Quote # {quote.quote_number} already exists. Enter a different quote number."
        }), 409
    return jsonify(quote_to_dict(quote))


@app.route("/quotes/<int:id>", methods=["DELETE"])
def delete_quote(id):
    quote = Quote.query.get_or_404(id)
    db.session.delete(quote)
    db.session.commit()
    return jsonify({"message": "Quote deleted"})

@app.route("/locations")
def get_locations():
    search = request.args.get("search", "").strip()

    query = (
        db.session.query(Quote.location)
        .filter(Quote.location.isnot(None))
        .filter(Quote.location != "")
    )

    if search:
        query = query.filter(Quote.location.ilike(f"%{search}%"))

    locations = query.distinct().limit(10).all()

    return jsonify([x[0] for x in locations])

# ---------------------------
# Line Items
# ---------------------------
@app.route("/quotes/<int:quote_id>/line-items", methods=["POST"])
def create_line_item(quote_id):
    Quote.query.get_or_404(quote_id)
    data = request.json or {}

    last_item = (
        LineItem.query.filter_by(quote_id=quote_id)
        .order_by(LineItem.sort_order.desc())
        .first()
    )
    next_sort_order = (last_item.sort_order + 1) if last_item else 1

    item = LineItem(
        quote_id=quote_id,
        sort_order=next_sort_order,
        tag=data.get("tag"),
        vendor=data.get("vendor"),
        qty=num(data.get("qty"), 1),
        description=data.get("description"),
        item=data.get("item"),
        type=data.get("type"),
        series=data.get("series"),
        model=data.get("model"),
        part_number=data.get("part_number"),
        list_price=num(data.get("list_price"), 0),
        multiplier=num(data.get("multiplier"), 1),
        markup=num(data.get("markup"), 0),
        freight=num(data.get("freight"), 0),
        startup=num(data.get("startup"), 0),
        surcharge=num(data.get("surcharge"), 0),
        terms=data.get("terms", "FFA"),
        notes=data.get("notes"),
        included=bool_value(data.get("included"), False),
    )

    calculate_line_item(item)
    db.session.add(item)
    db.session.commit()
    return jsonify(line_item_to_dict(item)), 201


@app.route("/line-items/<int:id>", methods=["PUT"])
def update_line_item(id):
    item = LineItem.query.get_or_404(id)
    data = request.json or {}

    for field in ["tag", "vendor", "description", "item", "type", "series", "model", "part_number", "terms", "notes"]:
        setattr(item, field, data.get(field, getattr(item, field)))

    item.qty = num(data.get("qty"), item.qty)
    item.list_price = num(data.get("list_price"), item.list_price)
    item.multiplier = num(data.get("multiplier"), item.multiplier)
    item.markup = num(data.get("markup"), item.markup)
    item.freight = num(data.get("freight"), item.freight)
    item.startup = num(data.get("startup"), item.startup)
    item.surcharge = num(data.get("surcharge"), item.surcharge)
    item.included = bool_value(data.get("included"), item.included)

    calculate_line_item(item)
    db.session.commit()
    return jsonify(line_item_to_dict(item))


@app.route("/line-items/<int:id>", methods=["DELETE"])
def delete_line_item(id):
    item = LineItem.query.get_or_404(id)
    db.session.delete(item)
    db.session.commit()
    return jsonify({"message": "Line item deleted"})


@app.route("/quotes/<int:quote_id>/line-items/reorder", methods=["PUT"])
def reorder_line_items(quote_id):
    Quote.query.get_or_404(quote_id)
    item_ids = (request.json or {}).get("item_ids", [])

    for index, item_id in enumerate(item_ids):
        item = LineItem.query.filter_by(id=item_id, quote_id=quote_id).first()
        if item:
            item.sort_order = index + 1

    db.session.commit()
    return jsonify(quote_to_dict(Quote.query.get_or_404(quote_id)))


@app.route("/line-items/<int:line_item_id>/notes", methods=["PUT"])
def replace_line_item_notes(line_item_id):
    item = LineItem.query.get_or_404(line_item_id)
    data = request.json or {}
    notes = data.get("notes", [])

    LineItemNote.query.filter_by(line_item_id=line_item_id).delete()

    for index, note_data in enumerate(notes):
        note = LineItemNote(
            line_item_id=item.id,
            note_library_id=note_data.get("note_library_id"),
            category=note_data.get("category"),
            label=note_data.get("label"),
            text=note_data.get("text"),
            is_custom=bool_value(note_data.get("is_custom"), False),
            is_selected=bool_value(note_data.get("is_selected"), True),
            sort_order=note_data.get("sort_order", index + 1),
        )
        db.session.add(note)

    db.session.commit()
    return jsonify(line_item_to_dict(item))


# ---------------------------
# Products
# ---------------------------
@app.route("/products", methods=["GET"])
def get_products():
    search = request.args.get("search", "").strip()
    query = Product.query
    if search:
        like = f"%{search}%"
        query = query.filter(
            or_(
                Product.name.ilike(like),
                Product.vendor.ilike(like),
                Product.manufacturer.ilike(like),
                Product.category.ilike(like),
                Product.type.ilike(like),
                Product.series.ilike(like),
                Product.model.ilike(like),
                Product.part_number.ilike(like),
                Product.description.ilike(like),
            )
        )
    return jsonify([product_to_dict(p) for p in query.order_by(Product.id.desc()).all()])


@app.route("/products", methods=["POST"])
def create_product():
    p = Product()
    data = request.json or {}
    for field in ["name", "tag", "vendor", "manufacturer", "category", "type", "series", "model", "part_number", "description", "notes"]:
        setattr(p, field, data.get(field))
    p.part_number = data.get("part_number") or None
    p.list_price = num(data.get("list_price"), 0)
    p.multiplier = num(data.get("multiplier"), 1)
    p.surcharge = num(data.get("surcharge"), 0)
    p.net_cost = num(data.get("net_cost"), 0)
    p.is_active = bool_value(data.get("is_active"), True)
    db.session.add(p)
    db.session.commit()
    return jsonify(product_to_dict(p)), 201


@app.route("/products/<int:id>", methods=["PUT"])
def update_product(id):
    p = Product.query.get_or_404(id)
    data = request.json or {}
    for field in ["name", "tag", "vendor", "manufacturer", "category", "type", "series", "model", "part_number", "description", "notes"]:
        setattr(p, field, data.get(field, getattr(p, field)))
    p.list_price = num(data.get("list_price"), p.list_price)
    p.multiplier = num(data.get("multiplier"), p.multiplier)
    p.surcharge = num(data.get("surcharge"), p.surcharge)
    p.net_cost = num(data.get("net_cost"), p.net_cost)
    p.is_active = bool_value(data.get("is_active"), p.is_active)
    db.session.commit()
    return jsonify(product_to_dict(p))


@app.route("/products/<int:id>", methods=["DELETE"])
def delete_product(id):
    p = Product.query.get_or_404(id)
    db.session.delete(p)
    db.session.commit()
    return jsonify({"message": "Product deleted"})


@app.route("/products/lookup", methods=["GET"])
def lookup_product():
    term = request.args.get("term", "").strip()
    if not term:
        return jsonify([])

    like = f"%{term}%"
    products = Product.query.filter(
        or_(
            Product.item.ilike(like) if hasattr(Product, "item") else Product.category.ilike(like),
            Product.name.ilike(like),
            Product.series.ilike(like),
            Product.model.ilike(like),
            Product.part_number.ilike(like),
        )
    ).limit(15).all()

    return jsonify([product_to_dict(p) for p in products])


# ---------------------------
# Notes Library
# ---------------------------
@app.route("/notes-library", methods=["GET"])
def get_notes():
    search = request.args.get("search", "").strip()
    item = request.args.get("item", "").strip()
    type_ = request.args.get("type", "").strip()
    series = request.args.get("series", "").strip()
    model = request.args.get("model", "").strip()
    note_type = request.args.get("note_type", "").strip()

    query = NotesLibrary.query

    if search:
        like = f"%{search}%"
        query = query.filter(
            or_(
                NotesLibrary.item.ilike(like),
                NotesLibrary.type.ilike(like),
                NotesLibrary.category.ilike(like),
                NotesLibrary.series.ilike(like),
                NotesLibrary.model.ilike(like),
                NotesLibrary.text.ilike(like),
                NotesLibrary.note_type.ilike(like),
            )
        )

    if item:
        query = query.filter(or_(NotesLibrary.item == item, NotesLibrary.item.is_(None)))
    if type_:
        query = query.filter(or_(NotesLibrary.type == type_, NotesLibrary.type.is_(None)))
    if series:
        query = query.filter(or_(NotesLibrary.series == series, NotesLibrary.series.is_(None)))
    if model:
        query = query.filter(or_(NotesLibrary.model == model, NotesLibrary.model.is_(None)))
    if note_type:
        query = query.filter(NotesLibrary.note_type == note_type)

    return jsonify([note_to_dict(n) for n in query.order_by(NotesLibrary.sort_order.asc(), NotesLibrary.id.desc()).all()])


@app.route("/notes-library", methods=["POST"])
def create_note():
    n = NotesLibrary()
    data = request.json or {}
    for field in ["item", "type", "category", "series", "model", "text", "note_type"]:
        setattr(n, field, data.get(field))
    n.default_selected = bool_value(data.get("default_selected"), False)
    n.sort_order = int(data.get("sort_order") or 0)
    n.is_active = bool_value(data.get("is_active"), True)
    db.session.add(n)
    db.session.commit()
    return jsonify(note_to_dict(n)), 201


@app.route("/notes-library/<int:id>", methods=["PUT"])
def update_note(id):
    n = NotesLibrary.query.get_or_404(id)
    data = request.json or {}
    for field in ["item", "type", "category", "series", "model", "text", "note_type"]:
        setattr(n, field, data.get(field, getattr(n, field)))
    n.default_selected = bool_value(data.get("default_selected"), n.default_selected)
    n.sort_order = int(data.get("sort_order") or n.sort_order or 0)
    n.is_active = bool_value(data.get("is_active"), n.is_active)
    db.session.commit()
    return jsonify(note_to_dict(n))


@app.route("/notes-library/<int:id>", methods=["DELETE"])
def delete_note(id):
    n = NotesLibrary.query.get_or_404(id)
    db.session.delete(n)
    db.session.commit()
    return jsonify({"message": "Note deleted"})


# ---------------------------
# Companies / Contacts
# ---------------------------
@app.route("/companies", methods=["GET"])
def get_companies():
    search = request.args.get("search", "").strip()
    query = Company.query
    if search:
        like = f"%{search}%"
        query = query.filter(or_(Company.name.ilike(like), Company.type.ilike(like), Company.city.ilike(like), Company.notes.ilike(like)))
    return jsonify([company_to_dict(c) for c in query.order_by(Company.name.asc()).all()])


@app.route("/companies", methods=["POST"])
def create_company():
    c = Company()
    data = request.json or {}
    for field in ["name", "type", "address1", "address2", "city", "state", "zipcode", "notes", "website", "account_number", "tax_id", "payment_terms"]:
        setattr(c, field, data.get(field))
    c.is_active = bool_value(data.get("is_active"), True)
    db.session.add(c)
    db.session.commit()
    return jsonify(company_to_dict(c)), 201


@app.route("/companies/<int:id>", methods=["PUT"])
def update_company(id):
    c = Company.query.get_or_404(id)
    data = request.json or {}
    for field in ["name", "type", "address1", "address2", "city", "state", "zipcode", "notes", "website", "account_number", "tax_id", "payment_terms"]:
        setattr(c, field, data.get(field, getattr(c, field)))
    c.is_active = bool_value(data.get("is_active"), c.is_active)
    db.session.commit()
    return jsonify(company_to_dict(c))


@app.route("/companies/<int:id>", methods=["DELETE"])
def delete_company(id):
    c = Company.query.get_or_404(id)
    db.session.delete(c)
    db.session.commit()
    return jsonify({"message": "Company deleted"})


@app.route("/contacts", methods=["GET"])
def get_contacts():
    search = request.args.get("search", "").strip()
    company_id = request.args.get("company_id")
    query = Contact.query
    if company_id:
        query = query.filter(Contact.company_id == int(company_id))
    if search:
        like = f"%{search}%"
        query = query.outerjoin(Company).filter(
            or_(
                Contact.first_name.ilike(like),
                Contact.last_name.ilike(like),
                Contact.email.ilike(like),
                Contact.role.ilike(like),
                Company.name.ilike(like),
            )
        )
    return jsonify([contact_to_dict(c) for c in query.order_by(Contact.first_name.asc()).all()])


@app.route("/contacts", methods=["POST"])
def create_contact():
    c = Contact()
    data = request.json or {}
    for field in ["company_id", "first_name", "last_name", "role", "email", "tel", "mobile", "notes"]:
        setattr(c, field, data.get(field))
    c.is_active = bool_value(data.get("is_active"), True)
    db.session.add(c)
    db.session.commit()
    return jsonify(contact_to_dict(c)), 201


@app.route("/contacts/<int:id>", methods=["PUT"])
def update_contact(id):
    c = Contact.query.get_or_404(id)
    data = request.json or {}
    for field in ["company_id", "first_name", "last_name", "role", "email", "tel", "mobile", "notes"]:
        setattr(c, field, data.get(field, getattr(c, field)))
    c.is_active = bool_value(data.get("is_active"), c.is_active)
    db.session.commit()
    return jsonify(contact_to_dict(c))


@app.route("/contacts/<int:id>", methods=["DELETE"])
def delete_contact(id):
    c = Contact.query.get_or_404(id)
    db.session.delete(c)
    db.session.commit()
    return jsonify({"message": "Contact deleted"})


@app.route("/auth/login", methods=["POST"])
def login():
    data = request.get_json() or {}

    email = data.get("email", "").strip().lower()
    password = data.get("password", "")

    user = User.query.filter_by(email=email, active=True).first()

    if not user or not check_password_hash(user.password_hash, password):
        return jsonify({"error": "Invalid email or password"}), 401

    session["user_id"] = user.id
    session["user_name"] = user.name
    session["user_email"] = user.email

    return jsonify({
        "id": user.id,
        "name": user.name,
        "email": user.email
    })


@app.route("/auth/me", methods=["GET"])
def auth_me():
    if "user_id" not in session:
        return jsonify({"user": None}), 401

    return jsonify({
        "user": {
            "id": session["user_id"],
            "name": session["user_name"],
            "email": session["user_email"]
        }
    })


@app.route("/auth/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"message": "Logged out"})




def format_contact(contact):
    if not contact:
        return ""

    # Already a list
    if isinstance(contact, list):
        return ", ".join(str(x).strip() for x in contact if x)

    # String that looks like a list: '["Mike Llorence"]' or '["A", "B"]'
    if isinstance(contact, str):
        contact = contact.strip()

        # Try to parse it as JSON
        try:
            parsed = json.loads(contact)
            if isinstance(parsed, list):
                return ", ".join(str(x).strip() for x in parsed if x)
        except Exception:
            pass

        # Fallback: remove [ ] and " characters
        cleaned = contact.replace("[", "").replace("]", "").replace('"', "").replace("'", "")
        return cleaned.strip()

    return str(contact)

def apply_quote_filters(query):
    """Reuse the same filters used by get_quotes"""
    search = request.args.get("search", "").strip()
    line_search = request.args.get("line_search", "").strip()
    status = request.args.get("status", "").strip()
    customer = request.args.get("customer", "").strip()
    location = request.args.get("location", "").strip()
    salesman = request.args.get("salesman", "").strip()
    quote_date_from = request.args.get("quote_date_from", "").strip()
    quote_date_to = request.args.get("quote_date_to", "").strip()
    bid_date_from = request.args.get("bid_date_from", "").strip()
    bid_date_to = request.args.get("bid_date_to", "").strip()

    if search:
        like = f"%{search}%"
        query = query.filter(
            or_(
                Quote.quote_number.ilike(like),
                Quote.project.ilike(like),
                Quote.attention.ilike(like),
            )
        )

    if line_search:
        like = f"%{line_search}%"
        query = query.join(LineItem).filter(LineItem.description.ilike(like))

    if status:
        query = query.filter(Quote.status == status)

    if customer:
        query = query.filter(Quote.to_company.ilike(f"%{customer}%"))

    if location:
        query = query.filter(Quote.location.ilike(f"%{location}%"))

    if salesman:
        query = query.filter(Quote.contact.cast(db.String).ilike(f"%{salesman}%"))

    if quote_date_from:
        query = query.filter(Quote.date >= datetime.strptime(quote_date_from, "%Y-%m-%d").date())

    if quote_date_to:
        query = query.filter(Quote.date <= datetime.strptime(quote_date_to, "%Y-%m-%d").date())

    if bid_date_from:
        query = query.filter(Quote.bid_date >= bid_date_from)

    if bid_date_to:
        query = query.filter(Quote.bid_date <= bid_date_to)

    return query.distinct()


def style_excel_worksheet(ws):
    """Make the Excel look clean"""
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
            cell.border = thin
        adjusted_width = min(max_length + 2, 45)
        ws.column_dimensions[column].width = adjusted_width


@app.route("/quotes/export", methods=["GET"])
def export_quotes():
    query = apply_quote_filters(Quote.query)
    quotes = query.order_by(Quote.id.desc()).all()

    rows = []
    for q in quotes:
        total = sum(float(item.total_price or 0) for item in q.line_items)
        # contact_str = ", ".join(q.contact) if isinstance(q.contact, list) else str(q.contact or "")
        contact_str = format_contact(q.contact)

        rows.append({
            "Quote #": q.quote_number,
            "Quote Date": q.date.strftime("%m/%d/%Y") if q.date else "",
            "Bid Due Date": q.bid_date or "",
            "Status": q.status or "",
            "Job / Project": q.project or "",
            "Customer": q.to_company or "",
            "Attn": q.attention or "",
            "Outside Sales": contact_str,
            "Location": q.location or "",
            "Total": round(total, 2),
            "Created By": q.created_by or "",
            "Notes": q.notes or "",
        })

        

    df = pd.DataFrame(rows)

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Quotes")
        ws = writer.sheets["Quotes"]
        style_excel_worksheet(ws)

    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"quotes_export_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    )


@app.route("/quotes/export-line-items", methods=["GET"])
def export_line_items():
    query = apply_quote_filters(Quote.query)
    quotes = query.order_by(Quote.id.desc()).all()

    rows = []
    for q in quotes:
        # contact_str = ", ".join(q.contact) if isinstance(q.contact, list) else str(q.contact or "")
        contact_str = format_contact(q.contact)

        for item in q.line_items:
            rows.append({
                "Quote #": q.quote_number,
                "Quote Date": q.date.strftime("%m/%d/%Y") if q.date else "",
                "Status": q.status or "",
                "Customer": q.to_company or "",
                "Job / Project": q.project or "",
                "Outside Sales": contact_str,
                "Tag": item.tag or "",
                "Qty": item.qty or 0,
                "Description": item.description or "",
                "Item / Category": item.item or "",
                "Type": item.type or "",
                "Series": item.series or "",
                "Model": item.model or "",
                "Part Number": item.part_number or "",
                "Vendor": item.vendor or "",
                "List Price": float(item.list_price or 0),
                "Multiplier": float(item.multiplier or 1),
                "Markup": float(item.markup or 0),
                "Freight": float(item.freight or 0),
                "Startup": float(item.startup or 0),
                "Surcharge": float(item.surcharge or 0),
                "Net Cost": float(item.net_cost or 0),
                "Sell Price": float(item.sell_price or 0),
                "Total Price": float(item.total_price or 0),
                "Terms": item.terms or "",
                "Included": "Yes" if item.included else "No",
            })

    df = pd.DataFrame(rows)

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Line Items")
        ws = writer.sheets["Line Items"]
        style_excel_worksheet(ws)

    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"line_items_export_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    )


if __name__ == "__main__":
    app.run(debug=True)
